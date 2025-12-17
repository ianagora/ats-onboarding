# public.py
from __future__ import annotations
import os
import importlib
from datetime import datetime
from uuid import uuid4
from typing import Optional

from flask import (
    Blueprint, render_template, request, redirect, url_for,
    session, flash, current_app, abort
)
from werkzeug.utils import secure_filename
from sqlalchemy.orm import Session as SASession
from sqlalchemy import select

public_bp = Blueprint(
    "public",
    __name__,
    template_folder="templates_public",
    static_folder="static_public",
    static_url_path="/public_static",
)

# ----------------- Helpers -----------------
def _models():
    """Lazy import to avoid circular import with app.py."""
    from app import engine, Application, Candidate, Job  # noqa: WPS433
    return engine, Application, Candidate, Job

def _load_from_app(name: str):
    """Safely load a callable/attr from app at runtime (avoid top-level imports)."""
    try:
        app_mod = importlib.import_module("app")
        return getattr(app_mod, name, None)
    except Exception:
        return None

def _get_user_id() -> Optional[int]:
    return session.get("public_user_id")

def _require_login(next_url: str):
    """If not logged in as a public user, redirect to public.login with ?next=..."""
    if not _get_user_id():
        return redirect(url_for("public.login", next=next_url))
    return None

def _static_upload_dir(*parts: str) -> str:
    root = current_app.root_path  # app.py folder
    return os.path.join(root, "static", *parts)

def _save_cv(file_storage) -> Optional[str]:
    """Save uploaded CV to /static/uploads/cvs and return relative path 'uploads/cvs/<file>'."""
    if not file_storage or not (file_storage.filename or "").strip():
        return None
    fname = secure_filename(file_storage.filename)
    ext = os.path.splitext(fname)[1].lower() or ".pdf"
    new_name = f"{uuid4().hex}{ext}"
    upload_dir = _static_upload_dir("uploads", "cvs")
    os.makedirs(upload_dir, exist_ok=True)
    file_storage.save(os.path.join(upload_dir, new_name))
    return f"uploads/cvs/{new_name}"

def _maybe_set(obj, name: str, value):
    """Set obj.name = value if the mapped attribute exists and value is truthy."""
    if value and hasattr(obj, name):
        try:
            setattr(obj, name, value)
        except Exception:
            pass

def _mirror_cv_to_candidate(cand, relpath: Optional[str]):
    """Mirror CV presence to Candidate for internal views."""
    if not cand or not relpath:
        return
    for attr in ("cv_path", "cv_url", "resume_path", "cv", "cv_filename", "cv_file_path"):
        if hasattr(cand, attr):
            _maybe_set(cand, attr, relpath)
            break
    if hasattr(cand, "has_cv"):
        try:
            cand.has_cv = True
        except Exception:
            pass

def _abs_from_static_rel(rel_path: str) -> str:
    root = current_app.root_path
    return os.path.join(root, "static", rel_path)

def _safe_extract_text(file_path: str, original_name: str) -> str:
    """Basic extractor for PDF/DOCX/TXT."""
    name = (original_name or file_path or "").lower()
    # PDF
    if name.endswith(".pdf"):
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                bits = [(p.extract_text() or "") for p in pdf.pages]
            txt = "\n".join(bits).strip()
            if txt:
                return txt
        except Exception:
            pass
    # DOCX
    if name.endswith(".docx"):
        try:
            import docx
            d = docx.Document(file_path)
            txt = "\n".join([p.text for p in d.paragraphs]).strip()
            if txt:
                return txt
        except Exception:
            pass
    # Fallback: TXT
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
    except Exception:
        return ""

def _job_text(job) -> str:
    for attr in ("description", "desc", "requirements", "summary", "details", "spec"):
        if hasattr(job, attr):
            v = getattr(job, attr) or ""
            v = v.strip()
            if v:
                return v
    return (getattr(job, "title", "") or "").strip()

def _first(d: dict, keys: list[str], default=None):
    """Return the first non-None value for any of the provided keys in dict d."""
    if not isinstance(d, dict):
        return default
    for k in keys:
        if k in d and d[k] is not None:
            return d[k]
    return default

def _normalise_score(raw) -> Optional[int]:
    """
    Accepts many shapes:
      - number (0–1 or 0–100)
      - "78" or "78%" strings
      - (score, explanation) tuple/list
    Returns an int 0..100, or None.
    """
    if raw is None:
        return None

    # Tuple like (score, explanation)
    if isinstance(raw, (tuple, list)) and raw:
        raw = raw[0]

    # Strings (possibly with %)
    if isinstance(raw, str):
        s = raw.strip()
        if s.endswith("%"):
            s = s[:-1].strip()
        try:
            val = float(s)
        except Exception:
            return None
    else:
        # Numbers
        try:
            val = float(raw)
        except Exception:
            return None

    # Scale if 0–1
    if 0.0 <= val <= 1.0:
        val *= 100.0

    # Clamp and round
    val = max(0.0, min(100.0, val))
    return int(round(val))

# ----------------- Auth -----------------
@public_bp.get("/auth/login")
def login():
    next_url = request.args.get("next") or url_for("public.jobs_index")
    return render_template("auth_login_public.html", next=next_url)

@public_bp.post("/auth/login")
def login_post():
    from app import engine, Candidate  # lazy import
    next_url = request.form.get("next") or request.args.get("next") or url_for("public.jobs_index")

    email = (request.form.get("email") or "").strip().lower()
    name = (request.form.get("name") or "").strip()

    if not email:
        flash("Please enter an email", "danger")
        return redirect(url_for("public.login", next=next_url))

    with SASession(engine) as s:
        cand = (
            s.query(Candidate)
            .filter(Candidate.email.ilike(email))
            .order_by(Candidate.id.desc())
            .first()
        )
        if not cand:
            cand = Candidate(name=name or email.split("@")[0], email=email)
            s.add(cand)
            s.commit()
        session["public_user_id"] = cand.id

    return redirect(next_url)

@public_bp.get("/auth/signup")
def signup():
    next_url = request.args.get("next") or url_for("public.jobs_index")
    return render_template("auth_signup_public.html", next=next_url)

@public_bp.post("/auth/signup")
def signup_post():
    from app import engine, Candidate  # lazy import
    next_url = request.form.get("next") or request.args.get("next") or url_for("public.jobs_index")

    email = (request.form.get("email") or "").strip().lower()
    name = (request.form.get("name") or "").strip()

    if not email or not name:
        flash("Please provide name and email", "danger")
        return redirect(url_for("public.signup", next=next_url))

    with SASession(engine) as s:
        exists = s.query(Candidate).filter(Candidate.email.ilike(email)).first()
        if exists:
            flash("Account exists — please log in.", "info")
            return redirect(url_for("public.login", next=next_url))
        cand = Candidate(name=name, email=email)
        s.add(cand)
        s.commit()
        session["public_user_id"] = cand.id

    return redirect(next_url)

@public_bp.get("/auth/logout")
def logout():
    session.pop("public_user_id", None)
    flash("Signed out", "success")
    return redirect(url_for("public.jobs_index"))

# ----------------- Public Jobs -----------------
@public_bp.get("/jobs")
def jobs_index():
    engine, _, _, Job = _models()
    with SASession(engine) as s:
        jobs = (
            s.query(Job)
             .filter((Job.status == "Open") | (Job.status == None))  # noqa: E711
             .order_by(Job.created_at.desc())
             .all()
        )
    return render_template("jobs_index.html", jobs=jobs)

@public_bp.get("/jobs/<int:job_id>")
def job_public(job_id: int):
    engine, _, Candidate, Job = _models()
    with SASession(engine) as s:
        job = s.get(Job, job_id)
        if not job or (job.status or "Open") != "Open":
            flash("Job not found or not open", "warning")
            return redirect(url_for("public.jobs_index"))
        cand = s.get(Candidate, _get_user_id()) if _get_user_id() else None
    return render_template("job_public.html", job=job, cand=cand)

# ----------------- Apply Flow -----------------
@public_bp.get("/jobs/<int:job_id>/apply")
def job_apply(job_id: int):
    need_login = _require_login(url_for("public.job_apply", job_id=job_id))
    if need_login:
        return need_login

    engine, _, Candidate, Job = _models()
    user_id = _get_user_id()

    with SASession(engine) as s:
        job = s.get(Job, job_id)
        if not job or (job.status or "Open") != "Open":
            flash("Job not found or not open", "warning")
            return redirect(url_for("public.jobs_index"))
        cand = s.get(Candidate, user_id) if user_id else None

    return render_template("apply_form_public.html", job=job, cand=cand)

@public_bp.post("/jobs/<int:job_id>/apply")
def job_apply_post(job_id: int):
    """
    Create/update an application, persist CV file, create a Document row for the CV,
    mirror CV to Candidate, then run AI summary + tagging + scoring immediately.
    """
    need_login = _require_login(url_for("public.job_apply", job_id=job_id))
    if need_login:
        return need_login

    engine, Application, Candidate, Job = _models()
    user_id = _get_user_id()

    cover_note_val = (request.form.get("cover_note") or "").strip()
    cv_file = request.files.get("cv_file")
    relpath = _save_cv(cv_file)
    original_name = (cv_file.filename if cv_file else "") or ""

    with SASession(engine) as s:
        job = s.get(Job, job_id)
        if not job or (job.status or "Open") != "Open":
            flash("Job not found or not open", "warning")
            return redirect(url_for("public.jobs_index"))

        cand = s.get(Candidate, user_id)
        if not cand:
            session.pop("public_user_id", None)
            return redirect(url_for("public.login", next=url_for("public.job_apply", job_id=job_id)))

        # Latest application per candidate/job
        existing = (
            s.query(Application)
             .filter(Application.candidate_id == cand.id, Application.job_id == job.id)
             .order_by(Application.created_at.desc())
             .first()
        )

        if existing:
            try:
                existing.status = existing.status or "Declared"
            except Exception:
                pass
            for attr in ("cv_file_path", "cv_path", "cv_url", "resume_path", "cv", "cv_filename"):
                _maybe_set(existing, attr, relpath)
            for attr in ("cover_note", "cover_letter", "note", "notes", "summary"):
                _maybe_set(existing, attr, cover_note_val)
            _mirror_cv_to_candidate(cand, relpath)
            s.commit()
            app_obj = existing
            app_id = existing.id
        else:
            app_obj = Application(
                candidate_id=cand.id,
                job_id=job.id,
                status="Declared",
                created_at=datetime.utcnow(),
            )
            s.add(app_obj)
            s.commit()
            for attr in ("cv_file_path", "cv_path", "cv_url", "resume_path", "cv", "cv_filename"):
                _maybe_set(app_obj, attr, relpath)
            for attr in ("cover_note", "cover_letter", "note", "notes", "summary"):
                _maybe_set(app_obj, attr, cover_note_val)
            _mirror_cv_to_candidate(cand, relpath)
            s.commit()
            app_id = app_obj.id

        # ---- Ensure a Document row for this CV ----
        Document = _load_from_app("Document")
        doc_row = None
        if Document and relpath:
            try:
                order_col = getattr(Document, "uploaded_at", Document.id)
                doc_row = s.execute(
                    select(Document)
                    .where(Document.candidate_id == cand.id)
                    .where(getattr(Document, "doc_type", Document.doc_type) == "cv")
                    .order_by(order_col.desc())
                    .limit(1)
                ).scalar_one_or_none()
                if not doc_row or (getattr(doc_row, "filename", None) != relpath):
                    payload = dict(candidate_id=cand.id, filename=relpath)
                    if hasattr(Document, "application_id"):
                        payload["application_id"] = app_id
                    if hasattr(Document, "doc_type"):
                        payload["doc_type"] = "cv"
                    if hasattr(Document, "original_name"):
                        payload["original_name"] = original_name
                    if hasattr(Document, "uploaded_at"):
                        payload["uploaded_at"] = datetime.utcnow()
                    doc_row = Document(**payload)
                    s.add(doc_row)
                    s.commit()
            except Exception as e:
                current_app.logger.exception("Failed to upsert Document row for CV: %s", e)

        # ---- AI: extract text -> summary -> tags/skills -> score ----
        extract_cv_text = _load_from_app("extract_cv_text")
        ai_summarise = _load_from_app("ai_summarise")
        ai_score_with_explanation = _load_from_app("ai_score_with_explanation")

        cv_text = ""
        if relpath:
            abs_path = _abs_from_static_rel(relpath)
            try:
                if callable(extract_cv_text):
                    try:
                        if doc_row is not None:
                            cv_text = extract_cv_text(doc_row) or ""
                            if not cv_text:
                                cv_text = extract_cv_text(abs_path) or ""
                        else:
                            cv_text = extract_cv_text(abs_path) or ""
                    except TypeError:
                        cv_text = extract_cv_text(abs_path) or ""
                if not cv_text:
                    cv_text = _safe_extract_text(abs_path, original_name)
            except Exception:
                cv_text = _safe_extract_text(abs_path, original_name)

        # 1) AI Summary
        summary = ""
        try:
            if callable(ai_summarise):
                summary = ai_summarise(cv_text or "") or ""
        except Exception as e:
            current_app.logger.exception("ai_summarise failed: %s", e)

        if hasattr(app_obj, "ai_summary"):
            try:
                app_obj.ai_summary = summary
            except Exception:
                pass
        if hasattr(cand, "ai_summary"):
            try:
                cand.ai_summary = summary
            except Exception:
                pass

        # 2) Tags/skills
        try:
            TaxonomyTag = _load_from_app("TaxonomyTag")
            CandidateTag = _load_from_app("CandidateTag")
            if TaxonomyTag and CandidateTag:
                text_lc = (cv_text or "").lower()
                matched_ids = set()
                all_tags = s.execute(select(TaxonomyTag)).scalars().all()
                for tg in all_tags:
                    token = (tg.tag or "").strip()
                    if token and token.lower() in text_lc:
                        matched_ids.add(tg.id)

                if matched_ids:
                    existing_ids = {
                        tid for (tid,) in s.execute(
                            select(CandidateTag.tag_id).where(CandidateTag.candidate_id == cand.id)
                        ).all()
                    }
                    for tid in matched_ids:
                        if tid not in existing_ids:
                            s.add(CandidateTag(candidate_id=cand.id, tag_id=tid))

                    # Mirror into Candidate.skills
                    all_tag_rows = s.execute(
                        select(TaxonomyTag)
                        .join(CandidateTag, CandidateTag.tag_id == TaxonomyTag.id)
                        .where(CandidateTag.candidate_id == cand.id)
                        .order_by(TaxonomyTag.tag.asc())
                    ).scalars().all()
                    tag_names = [t.tag for t in all_tag_rows if (t.tag or "").strip()]
                    cand.skills = ", ".join(dict.fromkeys(tag_names))
        except Exception as e:
            current_app.logger.debug("Tagging skipped: %s", e)

        # 3) AI Score vs Job (robust normalisation + multiple payload shapes)
        if (cv_text or "").strip():
            try:
                score_val: Optional[int] = None
                explanation_text: Optional[str] = None

                if callable(ai_score_with_explanation):
                    payload = ai_score_with_explanation(_job_text(job), cv_text)

                    if isinstance(payload, dict):
                        raw_score = _first(
                            payload,
                            ["blended_score", "score", "match_score", "overall", "total", "similarity", "similarity_score"],
                        )
                        score_val = _normalise_score(raw_score)
                        explanation_text = _first(
                            payload,
                            ["explanation", "reason", "rationale", "details", "why"],
                        )
                    elif isinstance(payload, (tuple, list)):
                        score_val = _normalise_score(payload[0] if payload else None)
                        if len(payload) > 1 and isinstance(payload[1], str):
                            explanation_text = payload[1]
                    else:
                        # plain number or string
                        score_val = _normalise_score(payload)

                # optional fallback scorer
                if score_val is None:
                    ai_score_simple = _load_from_app("ai_score_simple")
                    if callable(ai_score_simple):
                        score_val = _normalise_score(ai_score_simple(_job_text(job), cv_text))

                # persist score
                if score_val is not None:
                    if hasattr(app_obj, "ai_score"):
                        try:
                            app_obj.ai_score = score_val
                        except Exception:
                            pass
                    if hasattr(cand, "ai_score"):
                        try:
                            cand.ai_score = score_val
                        except Exception:
                            pass
                # persist explanation
                if explanation_text and hasattr(app_obj, "ai_explanation"):
                    try:
                        app_obj.ai_explanation = explanation_text
                    except Exception:
                        pass

            except Exception as e:
                current_app.logger.exception("AI scoring failed: %s", e)

        # 4) Ensure Shortlist (optional)
        Shortlist = _load_from_app("Shortlist")
        if Shortlist:
            try:
                exists_short = s.execute(
                    select(Shortlist).where(
                        Shortlist.candidate_id == cand.id,
                        Shortlist.job_id == job.id,
                    )
                ).scalar_one_or_none()
                if not exists_short:
                    payload = dict(candidate_id=cand.id, job_id=job.id)
                    if hasattr(Shortlist, "created_at"):
                        payload["created_at"] = datetime.utcnow()
                    s.add(Shortlist(**payload))
            except Exception:
                pass

        s.commit()

    # Optional legacy hook
    run_cv_scoring = _load_from_app("run_cv_scoring")
    try:
        if callable(run_cv_scoring):
            run_cv_scoring(app_id)
    except Exception:
        current_app.logger.exception("run_cv_scoring failed for app_id=%s", app_id)

    return redirect(url_for("public.apply_done", job_id=job_id, app_id=app_id))

@public_bp.get("/jobs/<int:job_id>/apply/done")
def apply_done(job_id: int):
    return render_template("apply_done_public.html", job_id=job_id)

# ----------------- Optional token-based route -----------------
@public_bp.get("/apply/<token>")
def public_apply(token: str):
    engine, _, _, Job = _models()
    with SASession(engine) as s:
        job = s.query(Job).filter(Job.public_token == token).first()
        if not job:
            abort(404)
        return redirect(url_for("public.job_apply", job_id=job.id))